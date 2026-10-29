"""Generate the isolated AIST Stakeholder Challenge dataset (Python 3.10+)."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile

import openpyxl

GENERATOR = "aist-stakeholder-challenge"
ROLE_COLUMNS = "Stakeholder_ID Stakeholder_Name Role_Description Key_Priorities AI_Adoption_Perspective Typical_Challenges Decision_Authority Role_Qualification_Question Status Source Blockers_Linked".split()
QUESTION_COLUMNS = "Question_ID Stage Question_Text Response_Code Response_Label Response_Description Display_Order Required Status".split()
# Versioned application rules: these are not columns in the source workbook.
RULES = {
    "Q01": {"response_type": "single_choice", "codes": {"YES", "SOMEWHAT", "NO"}},
    "Q02": {"response_type": "multiple_choice", "codes": {"RESPONSIBLE", "AFFECTED", "CONTRIBUTOR", "NOT_RELEVANT", "NOT_SURE"}},
    "Q03": {"response_type": "free_text", "codes": {"FREE_TEXT"}},
}


class ValidationError(ValueError):
    pass


def fail(where, message):
    raise ValidationError(f"{where}: {message}")


def text(value, where):
    if not isinstance(value, str) or not value.strip():
        fail(where, "expected non-empty text")
    return value.strip()


def integer(value, where, minimum=0):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not float(value).is_integer() or value < minimum:
        fail(where, f"expected an integer >= {minimum}")
    return int(value)


def read_rows(workbook, name, columns):
    if name not in workbook.sheetnames:
        fail(name, "missing worksheet")
    sheet = workbook[name]
    rows = sheet.iter_rows()
    header_cells = next(rows, ())
    headers = [c.value.strip() if isinstance(c.value, str) else c.value for c in header_cells]
    # Ignore trailing formatted, empty columns only.
    while headers and headers[-1] is None:
        headers.pop()
    if len(headers) != len(set(headers)) or None in headers:
        fail(name, "duplicate or blank column header")
    if set(headers) != set(columns):
        fail(name, f"column mismatch; missing={sorted(set(columns)-set(headers))}; unexpected={sorted(set(headers)-set(columns), key=str)}")
    records = []
    for cells in rows:
        if all(c.value is None for c in cells):
            continue
        for c in cells:
            if c.data_type in {"f", "e"}:
                fail(f"{name}!{c.coordinate}", "formulas and Excel errors are not accepted; use literal values")
        if any(c.value is not None for c in cells[len(headers):]):
            fail(name, "data found outside named columns")
        values = {h: cells[i].value for i, h in enumerate(headers)}
        records.append((cells[0].row, values))
    if not records:
        fail(name, "worksheet contains no records")
    return records


def build_dataset(source: Path):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        role_rows = read_rows(workbook, "Stakeholder_Roles", ROLE_COLUMNS)
        question_rows = read_rows(workbook, "Challenge_Questions", QUESTION_COLUMNS)
    finally:
        workbook.close()
    roles, questions, warnings = {}, {}, []
    for row_number, row in role_rows:
        where = f"Stakeholder_Roles!row {row_number}"
        role = {k: text(v, f"{where}/{k}") for k, v in row.items() if k != "Blockers_Linked"}
        sid = role.pop("Stakeholder_ID")
        if not re.fullmatch(r"STK-\d{2,}", sid):
            fail(where, f"invalid stakeholder ID {sid!r}")
        if sid in roles:
            fail(where, f"duplicate stakeholder ID {sid}")
        role["Blockers_Linked"] = integer(row["Blockers_Linked"], f"{where}/Blockers_Linked")
        # Original column names remain available, without lossy splitting of prose.
        roles[sid] = {"id": sid, **role}
    for row_number, row in question_rows:
        where = f"Challenge_Questions!row {row_number}"
        clean = {k: text(v, f"{where}/{k}") for k, v in row.items() if k != "Display_Order"}
        qid = clean["Question_ID"]
        if qid not in RULES:
            fail(where, f"unsupported question {qid!r}; explicitly extend RULES before adding a question")
        if clean["Required"].casefold() not in {"yes", "no"}:
            fail(where, "Required must be Yes or No")
        required = clean["Required"].casefold() == "yes"
        metadata = {"stage": clean["Stage"], "text": clean["Question_Text"], "required": required, "status": clean["Status"]}
        if qid not in questions:
            questions[qid] = {"id": qid, **metadata, "response_type": RULES[qid]["response_type"], "options": []}
        question = questions[qid]
        if any(question[k] != v for k, v in metadata.items()):
            fail(where, f"inconsistent repeated metadata for {qid}")
        code = clean["Response_Code"]
        if code not in RULES[qid]["codes"]:
            fail(where, f"unsupported response code {code!r} for {qid}")
        order = integer(row["Display_Order"], f"{where}/Display_Order", 1)
        if any(o["code"] == code or o["display_order"] == order for o in question["options"]):
            fail(where, f"duplicate response code or display order for {qid}")
        question["options"].append({"id": f"{qid}:{code}", "code": code, "label": clean["Response_Label"], "description": clean["Response_Description"], "display_order": order})
    if set(questions) != set(RULES):
        fail("Challenge_Questions", f"missing questions: {sorted(set(RULES)-set(questions))}")
    for qid, question in questions.items():
        if {o["code"] for o in question["options"]} != RULES[qid]["codes"]:
            fail(qid, "missing required response codes")
        question["options"].sort(key=lambda o: (o["display_order"], o["code"]))
    for sid, role in roles.items():
        if role["Role_Qualification_Question"] != questions["Q01"]["text"]:
            fail(sid, "Role_Qualification_Question differs from Q01")
        role["qualification_question_id"] = "Q01"
    questions["Q02"]["exclusive_response_codes"] = ["NOT_RELEVANT", "NOT_SURE"]
    questions["Q03"]["max_length"] = None
    if questions["Q03"]["required"] and "optional" in questions["Q03"]["stage"].casefold():
        warnings.append("Q03 is labelled Optional explanation but Required=Yes; source value preserved. Resolve in Excel before publishing.")
    if any(r["Status"].casefold() == "draft" for r in roles.values()) or any(q["status"].casefold() == "draft" for q in questions.values()):
        warnings.append("Dataset contains Draft content; generation does not imply editorial approval.")
    warnings.append("Q03 maximum text length is unspecified in the source; max_length=null means no limit is defined by this dataset.")
    payload = {
        "schema_version": "1.0.0",
        "generator": GENERATOR,
        "rules_version": "1.0.0",
        "stakeholders": [roles[k] for k in sorted(roles)],
        "questions": [questions[k] for k in sorted(questions)],
        "settings": {
            "qualification_question_id": "Q01",
            "blocker_question_id": "Q02",
            "explanation_question_id": "Q03",
            "blockers_source": "external_explorer_dataset",
            "qualification_actions": {"YES": "proceed", "SOMEWHAT": "proceed_with_limited_familiarity", "NO": "reassign_role"},
        },
        "warnings": warnings,
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
    payload["dataset_version"] = "sha256:" + hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    return payload


def write_dataset(payload, source, destination):
    # Constrain writes to this standalone artifact and reject foreign JSON files.
    if destination.name != "stakeholder-challenge.json":
        fail(destination, "output filename must be stakeholder-challenge.json")
    if destination.is_symlink() or destination.resolve() == source.resolve():
        fail(destination, "output cannot be a symlink or the source workbook")
    if destination.exists():
        previous = json.loads(destination.read_text(encoding="utf-8"))
        if not isinstance(previous, dict) or previous.get("generator") != GENERATOR:
            fail(destination, "refusing to replace a JSON file not owned by this generator")
    content = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n", dir=destination.parent, prefix=".challenge-", suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("stakeholder-challenge.json"))
    parser.add_argument("--check", action="store_true", help="Validate without writing")
    args = parser.parse_args(argv)
    try:
        payload = build_dataset(args.workbook)
        for warning in payload["warnings"]:
            print(f"WARNING: {warning}", file=sys.stderr)
        if not args.check:
            write_dataset(payload, args.workbook, args.output)
        print(f"OK: {len(payload['stakeholders'])} stakeholders, {len(payload['questions'])} questions; {payload['dataset_version']}")
        return 0
    except (ValidationError, OSError, ValueError, KeyError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
